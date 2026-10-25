"""Generate actual submission artifacts from the implemented graph and measurements.

Run from any folder: python tools/generate_deliverables.py
No online resources, screenshots fabricated as UI, or external rendering needed.
"""
import csv
import json
import math
import sys
from html import escape
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from campus.graph import Graph, load_campus
from campus.algorithms import bfs, dfs, dijkstra, floyd_warshall, prim, kruskal
from campus.learning import ALGORITHMS, COMPLEXITY_NOTE, PSEUDOCODE, CORRECTNESS
from reportlab.pdfgen import canvas
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, Table, TableStyle
from reportlab.lib.enums import TA_LEFT
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

DOCS=ROOT/'docs'
DOCS.mkdir(exist_ok=True)
graph,coordinates=load_campus()
route=dijkstra(graph,'Main Gate','Hostel')
prim_result=prim(graph,'Main Gate')
kruskal_result=kruskal(graph)
pairs=floyd_warshall(graph)
bench=list(csv.DictReader((DOCS/'benchmark_results.csv').open(encoding='utf-8')))
NAVY='#0b1220'; TEAL='#159b89'; INK='#172942'; MUTED='#56667b'; BLUE='#326fc9'


def generate_study_guides():
    entries=[
    ('Basics','What is a graph?','A collection of vertices and edges. Here vertices are campus locations and edges are walkable paths.'),
    ('Basics','What are V and E?','V counts vertices and E counts undirected edges once. Our original campus has V=10 and E=17.'),
    ('Basics','Why is this graph undirected and weighted?','Each path can be walked both ways at the same modelled distance. A numerical weight records metres.'),
    ('Basics','What is an adjacency list?','For each location, a dictionary stores its neighbors and distances. An undirected edge appears in two neighbor dictionaries.'),
    ('Basics','What is Big O?','An upper bound on resource growth as input size increases. O(V+E) grows with vertices plus edges; it is not an exact duration.'),
    ('Basics','What is the difference between time and space complexity?','Time counts how computational work grows. Space counts memory growth. State whether graph storage, results and tracing are included.'),
    ('Basics','How do a queue and stack differ?','A queue removes the oldest item (FIFO). A stack resumes the newest item (LIFO). BFS uses a queue; DFS uses stack frames.'),
    ('Algorithms','Why select BFS?','It makes layered exploration and connectivity easy to demonstrate and guarantees minimum hop distance on an unweighted graph.'),
    ('Algorithms','Why select DFS?','It demonstrates deep exploration and backtracking with a stack, complementing BFS on the same graph.'),
    ('Algorithms','Does BFS find the shortest walking route?','Not in general: it minimizes edge count, not weight. In the A-B=2, B-C=1, A-C=5 triangle it reaches C directly in one edge, although walking through B costs only 3.'),
    ('Algorithms','Why are BFS and DFS O(V+E)?','Each reachable vertex is processed once. Each undirected edge is examined from both endpoints. Two times E is still O(E).'),
    ('Algorithms','Why select Dijkstra?','Distances are nonnegative, and the campus is a small sparse weighted graph. A min-heap efficiently selects the next nearest unsettled vertex.'),
    ('Algorithms','What is relaxation?','Try candidate = distance[u] + weight(u,v). If it is smaller than distance[v], update v and its parent.'),
    ('Algorithms','Why does Dijkstra need nonnegative weights?','Its greedy proof settles the smallest tentative distance permanently. A later negative edge could improve a settled vertex and break that argument.'),
    ('Algorithms','Why use heapq?','It implements a binary min-heap. Push and pop cost logarithmic time. This project uses lazy entries and skips vertices already settled.'),
    ('Algorithms','What is Dijkstra complexity here?','O((V+E) log V) time for a simple graph and O(V+E) extra/result space because lazy heap entries may reach O(E). Tracing adds O(V+E) events.'),
    ('Algorithms','Why select Floyd-Warshall?','It exposes dynamic programming and produces one small distance matrix answering all pairs of campus queries, with next-hop reconstruction.'),
    ('Algorithms','What does the Floyd-Warshall recurrence mean?','d[i,j] = min(d[i,j], d[i,k]+d[k,j]). Compare the previous best route with one passing through the newly allowed intermediate k.'),
    ('Algorithms','Why is Floyd-Warshall O(V cubed)?','The three nested loops each iterate over all V vertices. Matrices require O(V squared) space; our optional comparison trace can need O(V cubed).'),
    ('Algorithms','Does Floyd-Warshall support negative edges?','In general it supports negative directed edges if there is no negative cycle, and negative diagonals can detect such cycles. This undirected campus model rejects negatives: a negative undirected edge creates a negative closed walk.'),
    ('Algorithms','What is a spanning tree?','A connected acyclic subgraph containing every vertex. A tree on V vertices has V-1 edges. A minimum spanning tree minimizes their total weight.'),
    ('Algorithms','Why select Prim?','It explains growing an economical connected campus network from a starting location by repeatedly choosing a cheapest crossing edge.'),
    ('Algorithms','Why select Kruskal?','It gives a contrasting edge-oriented MST strategy: sort all edges and accept only those joining different components.'),
    ('Algorithms','What is Disjoint Set Union?','A structure representing components. find returns a representative; union merges components. Equal representatives mean an edge would create a cycle.'),
    ('Algorithms','What do path compression and union by rank do?','Compression makes visited nodes point toward the root. Rank attaches the shallower root beneath the deeper root. Together operations take nearly constant amortized time, O(alpha(V)).'),
    ('Algorithms','Why are Prim and Kruskal correct?','The cut property says a cheapest edge crossing a cut is safe for some MST. Prim uses the visited cut; Kruskal uses cuts between current components.'),
    ('Algorithms','Must Prim and Kruskal choose identical edges?','No. They must obtain the same optimal total weight. Ties may produce distinct MSTs; this code uses deterministic tie rules.'),
    ('Algorithms','Why is Kruskal usually O(E log E)?','Sorting dominates. DSU initialization is O(V), so the disconnected-input bound is O(V + E log E), with near-linear union/find work.'),
    ('Project','What exact result should I demonstrate?','Main Gate → Academic Block → Library → Canteen → Hostel costs 300+150+180+200 = 830 m. Both MST algorithms return 1690 m with nine edges.'),
    ('Project','How is an MST different from a shortest route?','An MST minimizes the sum of its network edges. Dijkstra minimizes one source-to-destination journey. On this campus the tree journey Main Gate to Hostel is 830 m, but Main Gate to Auditorium along the MST is 510 m versus a 380 m shortest route.'),
    ('Project','Why does the route finder settle ten vertices?','This implementation completes all reachable single-source distances even when a target is selected. The node count represents the entire source search.'),
    ('Project','How are unreachable destinations represented?','Distances remain math.inf and the reconstructed path is empty. The UI says unreachable instead of drawing an invented route.'),
    ('Project','What happens on a disconnected graph?','A traversal reaches only its starting component. Prim restarts in other components; Kruskal naturally returns a minimum forest. Both explicitly report that no complete spanning tree exists.'),
    ('Project','What if source equals destination?','The route is just the one source vertex and its distance is zero. No edges are highlighted.'),
    ('Project','Are zero-weight edges allowed?','Yes, all six algorithms support them in this nonnegative model. They are useful for testing, though distinct campus sites realistically have positive walking distances.'),
    ('Project','How do path edits work?','Graph.set_edge validates and updates both directions. remove_edge deletes both directions. Session state retains changes across pages; Restore reloads the original JSON dataset.'),
    ('Project','Are the coordinates real?','No. They are fictional local diagram coordinates in metres. Route calculations use explicit edge weights, not a real map or straight-line heuristic.'),
    ('Project','How did you test correctness?','Known-answer tests check small graphs. All-pairs distances agree with independent Dijkstra runs. A tiny MST is checked by enumerating every V-1 edge subset and testing connectivity. UI tests exercise pages and editing.'),
    ('Project','Do measured timings prove Big O?','No. They are seven-run perf_counter medians on seeded graphs with tracing disabled. Machine load and Python overhead matter; theoretical growth follows from operation counts.'),
    ('Project','Why is this syllabus aligned?','It implements the six graph algorithms requested for Unit 4 and explains efficiency and Big O from Unit 1. Practical 8-11 mapping is user-provided; no separate official syllabus was supplied.'),
    ('Project','What are the major limitations?','Fictional static distances, symmetric paths, length-proportional cost, no accessibility or congestion model, session-only edits and small-graph visualization.'),
    ('Project','Where should I start explaining the code?','Start with data/campus.json and Graph.adj, then dijkstra: initialize, pop nearest, settle, relax, record parent, reconstruct. Explain a small test before showing the interface.')
    ]
    text=['# Viva preparation guide','Answers refer to the original fictional campus. Say distances in metres and distinguish a route from a network.']
    for i,(level,q,a) in enumerate(entries,1):
        text += [f'## {i}. [{level}] {q}',a]
    (DOCS/'VIVA_GUIDE.md').write_text('\n\n'.join(text)+'\n',encoding='utf-8')
    worked={
    'BFS':'Start A: queue [A]. Remove A; discover B then C, giving queue [B,C]. Remove B, then C. Order A,B,C. C is one hop away even though its direct edge weighs 5. Mark on enqueue so a second incoming edge does not duplicate C.',
    'DFS':'Start A and enter B, its alphabetically first neighbor. From B enter C. C has no unseen neighbor, so pop its frame; then pop B and A. Order A,B,C. The stack stores suspended neighbor iterators, just as recursive calls would.',
    'Dijkstra':'Initially d(A)=0, d(B)=infinity, d(C)=infinity. Settle A and propose B=2 and C=5. Pop B=2, improve C to 2+1=3 and parent[C]=B. Pop C=3 and settle it. The old heap entry C=5 is ignored. Reversing parents C,B,A gives A,B,C at distance 3.',
    'Floyd-Warshall':'Initially d(A,C)=5, d(A,B)=2, d(B,C)=1, with diagonal zero. When B is allowed as an intermediate, compare 5 with 2+1=3. Store d(A,C)=3 and next[A,C]=B. Follow A→B→C to reconstruct. The reverse pair also becomes 3 in this undirected graph.',
    'Prim':'Start A. The crossing heap offers A-B=2 and A-C=5. Accept A-B=2; B exposes B-C=1. Accept B-C=1. The total is 3 with two edges. Prim does not sort all graph edges globally; each chosen edge crosses its current visited cut.',
    'Kruskal':'Sort B-C=1, A-B=2, A-C=5. Initially {A},{B},{C}. Accept B-C and merge {B,C}. Accept A-B and merge all three. Reject A-C because its endpoints already have the same DSU root. Total 1+2=3.'}
    walkthrough={
    'BFS':'Read bfs in campus/algorithms.py. deque([source]) initializes FIFO storage; seen marks the source. popleft removes the oldest vertex. The neighbor loop marks unseen vertices before appending them. Result.order is the visit order; discover and visit events narrate the queue actions.',
    'DFS':'Read dfs. stack contains (vertex, neighbor_iterator) frames. next(neighbors, None) resumes the top frame; None means backtrack by pop. An unseen vertex is marked, appended to order, and pushed with its own iterator. No Python recursion is used, so a long chain does not hit the recursion limit.',
    'Dijkstra':'Read dijkstra. Result.distances starts at infinity except source. heappop selects a minimum pair; settled filters lazy duplicates. candidate adds one edge cost. A strict improvement updates the distance and parent and pushes a new pair. After the search, target is followed backwards through parent and the list is reversed. Omitting target requests only the complete source distance table.',
    'Floyd-Warshall':'Read floyd_warshall and AllPairs.path. Two nested dictionary comprehensions initialize distances and next_hop, then direct edges replace infinity. The k/u/v loops try each intermediate and record before, candidate and improvement fields for every comparison when tracing is enabled. A strict improvement copies next_hop[u][k] as the first hop toward v. AllPairs.path repeatedly follows those first hops; None means unreachable.',
    'Prim':'Read prim. roots starts with the requested source and then lists other vertices. An unseen root gets a (0, root, None) heap entry. Popping a new vertex accepts its parent edge, if present, then pushes outgoing edges toward unvisited vertices. The outer root loop handles disconnected components. len(edges)==V-1 and a nonempty graph determines the connected flag.',
    'Kruskal':'Read DisjointSet then kruskal. parent and rank are dictionaries. find first locates a root, then compresses the traversed path. union rejects equal roots and otherwise joins by rank. kruskal sorts (weight, endpoint names), accepts edges whose union succeeds, and records reject events for cycle edges.'}
    parts=['# Learn the project from zero',
    '## 1. Read this before opening the code\nA campus is a network. A **vertex** is a place, such as the Library. An **edge** is a direct walking path between two places. A **weighted graph** gives each edge a number: walking distance in metres. **Undirected** means both directions have the same weight. An adjacency list stores each place\'s neighbors. In this project it is a Python dictionary of dictionaries, graph.adj[u][v] = weight.',
    'A **path** is a sequence of adjacent locations. Its weighted length is the sum of its edge distances. A **cycle** returns to its starting point. A graph is **connected** when every pair is reachable. A **component** is one connected piece. A **tree** is connected and has no cycle; it has V-1 edges. A **forest** contains a tree for each component.',
    '## 2. Data structures and strategies\nA queue is like a line at the canteen: first in, first out. A stack is like a stack of trays: last in, first out. A min-heap retrieves the smallest item efficiently, without maintaining a completely sorted list. A greedy algorithm takes a locally best safe choice; correctness needs a proof, not merely intuition. Dijkstra and MST algorithms are greedy here; Floyd-Warshall is dynamic programming, solving overlapping subproblems by allowed intermediates.',
    '## 3. Set up and follow the data\nRun setup.bat and run.bat. The JSON dataset becomes a Graph in graph.py. A Streamlit selection calls a function in algorithms.py. That function returns a Result (or AllPairs) containing the answer and compact events. visualization.py draws that computed answer. app.py formats it for a reader. The logic never imports Streamlit, so tests run independently.',
    '## 4. One tiny graph for all six algorithms\nUse A-B=2, B-C=1, A-C=5. Draw a triangle and write the weights. This is a teaching example, separate from the single campus dataset. Neighbor order is alphabetical. Work through the six methods below with pencil and paper before trying the campus graph.']
    for name in PSEUDOCODE:
        row=next(r for r in ALGORITHMS if r['name']==name)
        parts += [f'## {name}: {row["objective"]}',f'**Hand-worked example.** {worked[name]}',f'**Pseudocode**\n```text\n{PSEUDOCODE[name]}\n```',f'**Walk through the Python.** {walkthrough[name]}',f'**Why it works.** {CORRECTNESS[name]}',f'**Efficiency.** {row["time"]} time; {row["space"]} auxiliary/result space with trace disabled. {row["advantage"]}. Limitation: {row["limitation"]}.']
    parts += ['## 5. Walk through the real campus\nThe original graph has 10 locations and 17 paths. Dijkstra from Main Gate first proposes Academic Block=300 and Administrative Block=140. It settles Administrative Block first; that gives Auditorium=380. Academic Block remains 300 because 140+160 ties the current distance, so the parent is unchanged. Academic Block proposes Library=450 and Laboratory=470. Auditorium proposes Canteen=640. Library later improves Canteen to 630 and Medical Centre to 730. Canteen improves Hostel to 830. The reconstructed route is Main Gate → Academic Block → Library → Canteen → Hostel: 300+150+180+200=830 m.',
    '**Prim selection from Main Gate:**\n'+ '\n'.join(f'{i}. {u} — {v}: {w} m' for i,(u,v,w) in enumerate(prim_result.edges,1)),
    f'Total = {prim_result.total} m. Kruskal reaches the same total but accepts edges in global nondecreasing weight order. The original campus has a unique MST despite some equal-weight edges; the complete equal-weight test graph demonstrates non-unique choices. Equal-weight shortest routes also exist: Main Gate to Academic Block can be direct (300) or via Administrative Block (140+160).',
    '## 6. Shortest route versus minimum network\nSuppose the university builds only the MST. Its total construction length is minimal. A journey inside that tree may still be longer: Main Gate to Auditorium follows Main Gate–Administrative Block–Academic Block–Auditorium for 140+160+210=510 m. The full graph has a shorter Main Gate–Administrative Block–Auditorium route of 380 m. Minimize one journey with Dijkstra; minimize total network length with an MST.',
    '## 7. Understand complexity without memorizing symbols\nFor BFS/DFS, count one processing step per visited vertex and two adjacency checks per undirected edge: V+2E becomes O(V+E). Floyd has three V loops: V×V×V. Heap pushes/pops cost logarithmic time; Dijkstra and lazy Prim do O(V+E) such work, giving O((V+E) log V) on simple graphs. Kruskal sorts E edges, O(E log E), and initializes V DSU entries. Big O omits constants and lower-order terms; it describes growth, not exact seconds. O(V²) space means that doubling V roughly quadruples table entries.',
    COMPLEXITY_NOTE,
    '## 8. Read and trust tests\nRun python -m unittest discover -s tests -v using your installed environment. A known-answer test asserts an independently worked result. The A-B-C-D graph has shortest A-to-D distance 6 and MST weight 6. A disconnected graph keeps infinity for unreachable pairs; its MST result is a forest. An exhaustive small-graph oracle enumerates candidate three-edge trees. Agreement between Dijkstra and Floyd across generated graphs checks more cases, but no finite set of tests is a universal proof.',
    '## 9. Common mistakes and memory cues\nBFS: “breadth before depth,” use a queue. DFS: “finish the branch,” use a stack. Dijkstra: “nearest unsettled first.” Floyd: “could k help i reach j?” Prim: “grow one connected patch.” Kruskal: “sort and join components.” Do not mix up a parent pointer for single-source reconstruction with the next-hop table for all-pairs reconstruction. Do not treat infinity as an actual path length or assume a forest is a spanning tree.',
    '## 10. A practical learning session\nFirst draw the triangle and calculate all answers manually. Next run the route finder and predict the next event before clicking Next. Then remove both gate edges and explain why the source is isolated. Restore the original graph, run the tests, and practise the five-minute script. Finish by explaining one relaxation and one DSU rejection without looking at the screen.']
    parts += ['## 11. How the repaired step viewer works\nThe fragment in ui/steps.py reruns only its controls and visualization. campus/playback.py keeps one execution per browser session and replays compact events into a Frame. A different page, algorithm, source, destination or actual graph fingerprint resets the index. Unchanged results and a base graph canvas are reused within that session. There is no global mutable graph cache. BFS distinguishes discovered nodes from visited nodes; DFS shows the active stack; Dijkstra shows tentative versus settled distances; Kruskal separates examining, accepting and rejecting an edge. Floyd records every I/J/K comparison and reconstructs a matrix from direct distances plus earlier improvements. Previous applies a shorter event prefix, so it restores exactly the earlier state. This takes O(V²+t) work for a Floyd seek to event t, with O(V²) replay tables instead of storing a full matrix at every step. The final summary matrix and the replay matrix are labelled separately.']
    (DOCS/'LEARN_THE_PROJECT.md').write_text('\n\n'.join(parts)+'\n',encoding='utf-8')


class Report:
    def __init__(self):
        self.c=canvas.Canvas(str(DOCS/'Project_Report.pdf'),pagesize=(595.28,841.89))
        self.c.setTitle('Smart Campus Navigator - DAA Project Report')
        self.c.setAuthor('[Student Name]')
        self.number=0
        self.y=0
        self.body=ParagraphStyle('body',fontName='Helvetica',fontSize=10.1,leading=14.4,textColor=colors.HexColor(INK),spaceAfter=8)

    def page(self,title,kicker='DESIGN AND ANALYSIS OF ALGORITHMS / 01AI0506'):
        if self.number:self.c.showPage()
        self.number+=1
        self.c.setFillColor(colors.HexColor(NAVY));self.c.rect(0,754,596,88,fill=1,stroke=0)
        self.c.setFillColor(colors.HexColor('#62ddc9'));self.c.setFont('Helvetica-Bold',8);self.c.drawString(42,811,kicker)
        self.c.setFillColor(colors.white);self.c.setFont('Helvetica-Bold',21);self.c.drawString(42,775,title)
        self.c.setFillColor(colors.HexColor(MUTED));self.c.setFont('Helvetica',8)
        self.c.drawString(42,27,'Smart Campus Navigator | Fictional campus | Submission report')
        self.c.drawRightString(553,27,f'{self.number:02d} / 12')
        self.y=726

    def p(self,text,bold=False):
        paragraph=Paragraph(('<b>'+escape(text)+'</b>') if bold else escape(text),self.body)
        _,h=paragraph.wrap(511,700)
        if self.y-h<49:raise ValueError(f'Report page {self.number} overflows: {text[:50]}')
        paragraph.drawOn(self.c,42,self.y-h);self.y-=h+9

    def heading(self,text):
        self.y-=7
        self.c.setFillColor(colors.HexColor(TEAL));self.c.setFont('Helvetica-Bold',12)
        self.c.drawString(42,self.y,text);self.y-=21

    def code(self,text):
        lines=text.splitlines();height=13*len(lines)+16
        if self.y-height<49:raise ValueError(f'Code overflows page {self.number}')
        self.c.setFillColor(colors.HexColor('#edf3f8'));self.c.roundRect(42,self.y-height,511,height,5,fill=1,stroke=0)
        self.c.setFillColor(colors.HexColor(INK));self.c.setFont('Courier',8.4)
        for i,line in enumerate(lines):self.c.drawString(52,self.y-17-13*i,line)
        self.y-=height+12

    def table(self,rows,widths):
        style=ParagraphStyle('cell',parent=self.body,fontSize=8.1,leading=10.6)
        cells=[[Paragraph(escape(str(x)),style) for x in row] for row in rows]
        table=Table(cells,colWidths=widths,repeatRows=1)
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#d9ebe9')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f1f5f9')]),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]))
        _,height=table.wrap(511,700)
        if self.y-height<49:raise ValueError(f'Table overflows page {self.number}')
        table.drawOn(self.c,42,self.y-height);self.y-=height+12

    def graph(self,edges=(),height=285):
        x0=78;y0=self.y-height+20;s=.65
        selected={frozenset(e[:2]) for e in edges}
        for u,v,w in graph.edges():
            x1,y1=coordinates[u];x2,y2=coordinates[v]
            self.c.setStrokeColor(colors.HexColor(TEAL if frozenset((u,v)) in selected else '#b3c1d2'))
            self.c.setLineWidth(2.8 if frozenset((u,v)) in selected else 1)
            self.c.line(x0+s*x1,y0+.48*y1,x0+s*x2,y0+.48*y2)
            self.c.setFillColor(colors.HexColor(MUTED));self.c.setFont('Helvetica',6.4)
            self.c.drawCentredString(x0+s*(x1+x2)/2,y0+.48*(y1+y2)/2+3,str(w))
        for n,(x,y) in coordinates.items():
            self.c.setFillColor(colors.HexColor(BLUE));self.c.circle(x0+s*x,y0+.48*y,4,fill=1,stroke=0)
            self.c.setFillColor(colors.HexColor(INK));self.c.setFont('Helvetica-Bold',7)
            self.c.drawCentredString(x0+s*x,y0+.48*y+9,n)
        self.y-=height


def generate_report():
    r=Report()
    r.page('Smart Campus Navigator','B.TECH / ARTIFICIAL INTELLIGENCE AND MACHINE LEARNING')
    r.y=701
    r.p('Interactive Graph Algorithm Visualizer and Route Optimizer',True)
    r.p('Design and Analysis of Algorithms mini-project | Course code: 01AI0506')
    r.graph(list(zip(route.path,route.path[1:])),280)
    r.p('A fictional ten-location campus connects shortest-path learning with economical network planning. All six algorithms are implemented manually in Python and explained through inspectable events.')
    r.heading('Student submission details')
    for label,key in [('Student name','student'),('Enrollment number','enrollment'),('Guide name','guide'),('University','university')]:
        r.c.setFillColor(colors.HexColor(INK));r.c.setFont('Helvetica',10);r.c.drawString(42,r.y-15,label)
        r.c.acroForm.textfield(name=key,tooltip=label,x=181,y=r.y-22,width=368,height=24,value=f'[{label}]',borderWidth=1,borderColor=colors.HexColor('#c1ccda'),fillColor=colors.HexColor('#f2f6fa'),textColor=colors.HexColor(INK),fontSize=10)
        r.y-=38
    r.p('Editable fields above can be filled in a compatible PDF reader. All campus coordinates and distances are illustrative, not a real university survey.')

    r.page('Abstract and introduction')
    r.heading('Abstract')
    r.p('Smart Campus Navigator models a fictional university as a weighted undirected graph and makes six fundamental algorithms visible. It addresses two distinct tasks: students seek short walking routes, while administrators seek a connected pathway network with minimum total length. BFS and DFS support exploration, Dijkstra finds single-source weighted shortest paths, Floyd-Warshall computes all-pairs distances, and Prim and Kruskal construct minimum spanning trees or forests.')
    r.p('The project uses Python 3.11 or newer, Streamlit for controls, and Plotly for an illustrative campus diagram. The algorithms rely on standard-library data structures and handwritten logic rather than graph-solving packages. Compact event histories expose the choices made during execution. A session-level editor permits controlled experiments with distances and connectivity.')
    r.p(f'On the original dataset, the shortest Main Gate-to-Hostel route is {route.distance} metres. Both MST methods produce a total construction length of {prim_result.total} metres across nine selected edges. These outputs are computed from the actual dataset and checked by automated tests. Repeated timing measurements are recorded separately and interpreted as machine-dependent experiments, not as proofs of asymptotic complexity.')
    r.heading('Introduction')
    r.p('A graph is a natural representation whenever places are linked by available movement. Straight-line proximity alone is insufficient: a student can only use listed walkable paths, and an indirect route may cost less than another available connection. The adjacency-list model records exactly those connections and their walking distances.')
    r.p('A visual interface helps bridge pseudocode and execution. A learner can inspect a Dijkstra relaxation, compare traversal orders, and see why a cycle edge is rejected during Kruskal. The same fixed network supports all examples, reducing the burden of learning a new dataset for each technique.')
    r.heading('Project contribution')
    r.p('The work combines readable independent algorithm functions, deterministic graph editing, a local interactive dashboard, known-answer testing and beginner-oriented learning materials. The contribution is an educational integration of established algorithms; it does not claim a new routing algorithm or a real campus deployment.')

    r.page('Problem, objectives and syllabus')
    r.heading('Problem statement')
    r.p('Given ten campus locations and seventeen bidirectional paths with finite nonnegative distances, compute useful routes and explain network structure. The application must answer a selected journey, analyze all location pairs, demonstrate reachable exploration, and construct a connected network using the least total pathway length when a spanning tree exists.')
    r.heading('Objectives')
    for text in ['Represent one fictional campus consistently in code, diagrams and examples.', 'Implement BFS, DFS, heap-based Dijkstra, Floyd-Warshall with next hops, heap-based Prim, and Kruskal with Disjoint Set Union.', 'Explain individual decisions through compact step events and visible highlights.', 'Support path changes and clearly distinguish unreachable routes and disconnected forests.', 'Relate operation counts to Big O and verify outcomes with automated tests.']:r.p('• '+text)
    r.heading('Syllabus relevance')
    r.p('The supplied brief identifies course 01AI0506, Unit 4 (Greedy Method and Graph Algorithms), Unit 1 support, and Practicals 8-11. No separate official syllabus document was included. This report therefore uses the requested mapping and does not assign invented official practical titles. Dijkstra, Prim and Kruskal demonstrate greedy choice; BFS and DFS demonstrate graph traversal; Floyd-Warshall demonstrates dynamic programming for all-pairs routes.')
    r.p('Unit 1 support appears in input-size definitions, time and space bounds, efficiency discussion, and measured performance. A viva explanation should distinguish theoretical analysis from observed runtime. Exact practical numbering should be reconciled with the department copy if it differs from the brief.')
    r.heading('Existing approach and proposed solution')
    r.p('A static campus sketch shows paths but requires manual route addition. A spreadsheet can hold distances but does not reveal the search process. The proposed system links computations to a diagram and readable events, allowing a learner to change an edge and explain the resulting answer. It stays local and small enough for an examiner to inspect the complete logic.')

    r.page('Graph model and system architecture')
    r.p('Vertices: Main Gate, Administrative Block, Academic Block, Library, Laboratory, Auditorium, Canteen, Hostel, Sports Ground and Medical Centre. An edge stores one walking distance in metres and appears in both endpoint adjacency dictionaries. The dataset is simple: no self loops and at most one edge per unordered pair.')
    r.graph(height=285)
    r.code('data/campus.json -> Graph validation / session edits\n                         |\n                 algorithms.py (six methods)\n                         |\n                 Result / AllPairs + events\n                         |\n                 app.py + visualization.py\n                         |\n                 Streamlit controls + Plotly')
    r.p('Coordinates are local illustrative positions in metres. Original path lengths are at least their endpoint straight-line separations; they may be longer because walking paths bend. Computations use edge weights, not coordinate geometry. User edits may deliberately be unrealistic to expose algorithm behavior.')
    r.p('The graph layer owns validation. The algorithm layer owns answers and history. The interface owns selection, session state and rendering. Tests call the same public functions used by the app. Report and slide generation import the original dataset and recorded measurements, avoiding a second hand-maintained source of sample answers.')

    r.page('Traversal: BFS and DFS')
    r.heading('Breadth-First Search')
    r.p('BFS uses a FIFO queue. A vertex is marked when enqueued so multiple incoming edges do not queue it twice. Removing the oldest queued vertex explores the graph in layers. In an unweighted graph those layers correspond to minimum edge counts, not minimum weighted distance.')
    r.code(PSEUDOCODE['BFS'])
    r.p('Triangle example A-B=2, B-C=1, A-C=5: from A the queue becomes [B,C], then each is removed. C is one hop away through the costly direct edge; this is why a weighted route finder uses Dijkstra.')
    r.heading('Depth-First Search')
    r.p('DFS explores a branch until it has no unseen neighbor, then backtracks. This implementation uses a stack of suspended neighbor iterators to reproduce recursive DFS without a recursion-depth limit. Every reachable vertex is entered once, and every neighbor list is completely scanned.')
    r.code(PSEUDOCODE['DFS'])
    r.p('On the same triangle DFS enters A, then B, then C, before popping exhausted frames. The visit order happens to match BFS on three vertices, but the exploration mechanism differs. On the campus DFS enters deep hostel and medical branches before returning to Sports Ground.')
    r.p('Both methods run in O(V+E) time with adjacency lists: process each vertex once and inspect each undirected edge twice. Queue/stack, seen set and returned order occupy O(V) space. A disconnected graph returns only the starting component; the UI reports the reached count.')

    r.page('Dijkstra: shortest walking routes')
    r.p('Dijkstra maintains the best known distance to each location, starting with zero at the source and infinity elsewhere. A binary min-heap chooses a smallest tentative vertex. Once settled, a vertex is final because nonnegative remaining edges cannot create a cheaper route through an unsettled vertex. Relaxation proposes improved distances through the settled location.')
    r.code(PSEUDOCODE['Dijkstra'])
    r.heading('Worked campus relaxation')
    r.p('Main Gate proposes Academic Block=300 and Administrative Block=140. Administrative Block is settled first and proposes Auditorium=380. Its alternative Academic Block route is also 300, so the strict improvement test keeps the existing direct parent. Academic Block proposes Library=450 and Laboratory=470. Auditorium first proposes Canteen=640; Library later improves it to 630. Canteen gives Hostel=830.')
    r.p('Following parents backward from Hostel gives Hostel, Canteen, Library, Academic Block, Main Gate. Reversing yields the displayed route, with 300+150+180+200=830 metres. The route finder completes all reachable single-source distances; its ten settled nodes reflect that full computation.')
    r.heading('Correctness and implementation assumptions')
    r.p(CORRECTNESS['Dijkstra'])
    r.p('The graph rejects negative or nonfinite weights and permits zero-weight edges. Lazy heap entries avoid decrease-key code; stale settled entries are skipped. A source equal to its destination yields [source] at zero distance. An unreachable destination keeps infinity and returns an empty path.')
    r.p('For a simple graph the binary-heap implementation takes O((V+E) log V) time. The lazy heap may retain O(E) entries, giving O(V+E) auxiliary/result memory, excluding the graph. Parent reconstruction uses O(V) time and output space in the longest simple route.')

    r.page('Floyd-Warshall: every pair')
    r.p('Floyd-Warshall is useful when many source-destination questions are asked on a small unchanged graph. It starts with zero on the diagonal, edge weights for direct connections and infinity for missing connections. A next-hop table supports reconstruction instead of storing only distances.')
    r.code(PSEUDOCODE['Floyd-Warshall'])
    r.heading('Intermediate-node invariant')
    r.p(CORRECTNESS['Floyd-Warshall'])
    r.p('For A-B=2, B-C=1 and A-C=5, allowing B as an intermediate improves d[A,C] to 3. The update stores next[A,C]=B. Reconstruction begins at A, follows B, then reaches C. The same reasoning applies to every campus pair and the all-pairs table reports 830 metres from Main Gate to Hostel.')
    r.heading('Selected computed distance entries')
    selected=[('Main Gate','Library'),('Main Gate','Auditorium'),('Main Gate','Hostel'),('Library','Medical Centre'),('Hostel','Sports Ground')]
    r.table([['From','To','Shortest distance (m)']]+[[u,v,pairs.distances[u][v]] for u,v in selected],[180,180,151])
    r.p('The outer intermediate loop and two endpoint loops each run V times, so time is O(V³). The two V-by-V tables use O(V²) space with history disabled. Compact update events may number O(V³), an additional tracing cost explicitly separated in the complexity discussion.')
    r.p('The general algorithm can handle negative directed edges without negative cycles, but this project uses the shared nonnegative undirected graph contract. An unreachable pair has no next hop. Selecting the same source and destination reconstructs a one-vertex zero-distance route.')

    r.page('Minimum network: Prim and Kruskal')
    r.p('An MST includes all vertices, has no cycle, and minimizes the sum of selected edge weights. With construction cost assumed proportional to length, it is a simple model of economical campus connectivity. It does not minimize all individual journeys. For example, Main Gate-to-Auditorium is 510 m in the MST but 380 m in the full graph.')
    r.code(PSEUDOCODE['Prim'])
    r.code(PSEUDOCODE['Kruskal'])
    r.p('Prim grows a component by taking a cheapest crossing edge from a min-heap. Kruskal considers globally sorted edges; DSU determines whether they join separate components. Path compression and union by rank keep repeated component queries efficient. Both follow the cut property: a cheapest edge across a cut is safe for some MST.')
    r.table([['Method','Original campus result','Selection rule'],['Prim',f'{prim_result.total} m; 9 edges','Cheapest crossing edge'],['Kruskal',f'{kruskal_result.total} m; 9 edges','Sorted edges; reject cycles']],[90,170,251])
    r.p('Prim takes O((V+E) log V) time and O(V+E) auxiliary/result space with its lazy edge heap. Kruskal takes O(V+E log E), usually written O(E log E) for connected graphs, and O(V+E) space. If input is disconnected, the code returns an optimal forest across components and explicitly says that a full spanning tree cannot be constructed.')
    r.p('Equal weights can permit different optimal MSTs, all with the same total. This campus happens to have a unique MST; a complete four-vertex equal-weight test graph exercises ties. Deterministic heap and sort keys keep live demonstrations reproducible.')

    r.page('Implementation and verification method')
    r.heading('Implementation methodology')
    r.p('1. Define and validate one JSON dataset. Graph.set_edge updates both directions, rejects unknown endpoints, self loops and invalid distances, and sorts affected neighbor lists at edit time. Graph.remove_edge removes both directions. Reloading the JSON restores the original.')
    r.p('2. Implement algorithms in independent functions. Result contains traversal order, route, distances, tree edges, total, connectivity status and events as applicable. AllPairs stores distance and next-hop tables. trace=False suppresses history storage for fairer timing.')
    r.p('3. Build Streamlit sections around those public functions. Plotly highlights only supplied result/step edges. Session state keeps one active execution and step index. A graph fingerprint invalidates it after edits. A Streamlit fragment reruns only playback controls and visual state. Previous, Next and Reset are explicit manual controls, avoiding fragile timer behavior during a viva.')
    r.p('4. Check behavior before documenting it. Automated tests use small known-answer graphs, validate exceptions and edge cases, compare Dijkstra against Floyd-Warshall, and enumerate tiny candidate trees as an independent MST oracle. Streamlit integration tests exercise all six pages and form callbacks.')
    r.heading('Representative actual checks')
    r.table([['Case','Expected and tested behavior'],['A-B=2, B-C=1, A-C=5, C-D=3, B-D=8','A-to-D = 6 via A,B,C,D; MST total = 6'],['Zero-weight triangle','A-to-C = 0; both MST totals = 0'],['Disconnected A-B=2, C-D=1, isolated E','A-to-E unreachable; forest total 3; connected=False'],['Same source and destination','Single-vertex path; distance 0'],['Campus shortcut edit','Main Gate-Hostel=10 makes route 10; restore returns 830'],['Long chain','Iterative DFS visits 1200 vertices without recursion']],[170,341])
    r.p('The committed test log records the executed unittest results. Finite tests support confidence but do not replace the invariants and cut-property arguments described earlier. Windows helpers are supplied for convenience; runtime checks in this workspace use Linux and Python 3.12.')

    r.page('Complexity and efficiency analysis')
    r.p('Big O is an upper bound on growth, suppressing constants and smaller terms. Here V means locations and E means undirected paths. An adjacency list stores each edge twice but still needs O(V+E) space. The table excludes graph storage and optional event history; it includes auxiliary structures and returned results.')
    r.table([['Algorithm','Time','Space','Main reason']]+[[a['name'],a['time'],a['space'],a['strategy']] for a in ALGORITHMS],[91,130,91,199])
    r.heading('Why these bounds fit this code')
    r.p('BFS and DFS process vertices once and scan each adjacency entry once. The queue, stack frames, visited set and order need O(V). Floyd-Warshall visits V intermediate choices times V source choices times V target choices; the distance and next-hop tables each have V² entries.')
    r.p('Dijkstra and Prim use lazy binary heaps, which may retain O(E) entries rather than only one entry per vertex. For this simple graph E is at most V(V-1)/2, so log E is O(log V). At most O(V+E) processing/push/pop work yields O((V+E) log V), with O(V+E) memory.')
    r.p('Kruskal initializes V DSU elements, sorts E edges in O(E log E), and performs O(E) union/find operations with inverse-Ackermann amortized cost. For nontrivial simple graphs sorting dominates that near-linear DSU work. The explicit O(V) term remains necessary when isolated vertices are present.')
    r.heading('Trace and preparation overhead')
    r.p('Histories store compact event deltas rather than a full copy of every distance table. Floyd records each comparison; replay rebuilds the evolving matrix for backward or forward navigation. BFS/DFS/Dijkstra/Prim/Kruskal add O(V+E) events, while Floyd may add O(V³). Event strings and dictionary allocation increase measured constants. The graph sorts neighbor dictionaries during editing; that preparation is outside algorithm and timing bounds.')
    r.p('Comparisons should match objectives: BFS versus DFS for exploration and Prim versus Kruskal for network construction. Dijkstra computes one source; Floyd-Warshall computes every pair. A faster single-source timing does not establish that it is universally better for repeated all-pairs workloads.')

    r.page('Demonstration and measured results')
    r.heading('Computed campus outputs')
    r.p('Shortest Main Gate-to-Hostel route: Main Gate - Academic Block - Library - Canteen - Hostel. Distance: 830 m. MST total: 1690 m with nine edges. The same source/target distance appears in the Floyd-Warshall table. These values are generated from the current original JSON dataset and asserted in tests.')
    r.p('Prim accepts (metres): Gate-Administrative 140; Administrative-Academic 160; Academic-Library 150; Academic-Laboratory 170; Library-Canteen 180; Canteen-Hostel 200; Academic-Auditorium 210; Auditorium-Sports 210; Hostel-Medical 270. The displayed application uses the full location names.')
    r.heading('Experimental timings: actual recorded run')
    environment=(DOCS/'benchmark_environment.txt').read_text().splitlines()
    r.p(' / '.join(environment[:2])+'. Each workload receives one warm-up followed by seven perf_counter measurements. The table reports median milliseconds. Graphs are a guaranteed connected chain plus seeded extra edges with probability 0.12. Input generation and explanatory history are excluded.')
    timing=[['Algorithm','V=12, E=17','V=24, E=51','V=48, E=165']]
    for a in ALGORITHMS:
        timing.append([a['name']]+[f'{float(next(x["Median ms"] for x in bench if x["Algorithm"]==a["name"] and int(x["V"])==n)):.4f}' for n in (12,24,48)])
    r.table(timing,[133,126,126,126])
    r.p('Small timings are strongly influenced by interpreter overhead and machine load. As graph size increases, both vertex count and edge density change. These observations illustrate this run; they neither prove Big O nor provide an objective-neutral ranking. Re-run campus.benchmark to obtain measurements for another machine.')
    r.heading('Live demonstration sequence')
    r.p('Start at Dashboard, show the 830-m route, compare BFS/DFS, inspect both 1690-m MSTs, and check the all-pairs matrix. Add a temporary shortcut to demonstrate recomputation, then restore the original graph. docs/DEMO_SCRIPT.md gives a timed five-minute sequence with exact control labels.')

    r.page('Conclusion, scope and references')
    r.heading('Advantages and limitations')
    r.p('The project keeps established algorithms visible in short readable functions, reuses one dataset and explains results with deterministic steps. It requires no account, key, database or external map service. After dependencies are installed, local execution supports an offline classroom demonstration.')
    r.p('The campus is fictional, distances are static and symmetric, and construction cost is assumed proportional to length. The model omits accessibility, crowding, slopes, permissions and turn restrictions. A minimum tree also has no redundant links. Session changes are temporary. Floyd-Warshall and extensive event histories suit small teaching graphs rather than city-scale road networks.')
    r.heading('Future scope')
    r.p('Useful extensions would collect validated real campus data with permission, mark accessibility constraints, add travel-time scenarios and persist reviewed graph versions. Each would need a clear data contract and new tests. The present submission deliberately avoids unrelated algorithms and external-service dependencies.')
    r.heading('Conclusion')
    r.p('Smart Campus Navigator turns six syllabus algorithms into an inspectable local demonstration. The tested campus route is 830 m, while the optimal network spans all ten locations with 1690 m of pathway. The distinction between those objectives, together with invariants, complexity analysis and edge-case handling, provides the main academic learning outcome.')
    r.heading('References')
    references=[
    'Cormen, Leiserson, Rivest and Stein. Introduction to Algorithms, 4th edition. MIT Press, 2022. Sections on traversal, shortest paths and minimum spanning trees.',
    'Python documentation: collections.deque, heapq, time.perf_counter and unittest. https://docs.python.org/3/library/',
    'Streamlit documentation. https://docs.streamlit.io/',
    'Plotly Python documentation. https://plotly.com/python/',
    'ReportLab User Guide. https://docs.reportlab.com/reportlab/userguide/',
    'python-pptx documentation. https://python-pptx.readthedocs.io/',
    'User-provided project brief: course 01AI0506; Unit 4 focus, Unit 1 support and Practical 8-11 mapping. A separate official syllabus was not attached.'
    ]
    for i,reference in enumerate(references,1):r.p(f'[{i}] {reference}')
    r.c.save()
    assert r.number==12


def generate_slides():
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    def rgb(hex_value):return RGBColor.from_string(hex_value.lstrip('#'))
    def text(slide,x,y,w,h,value,size=22,color='#e9f0fb',bold=False):
        box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        frame=box.text_frame;frame.word_wrap=True
        for i,line in enumerate(value.split('\n')):
            p=frame.paragraphs[0] if i==0 else frame.add_paragraph();p.text=line
            p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
            p.space_after=Pt(12)
        return box
    def slide(title,kicker='SMART CAMPUS / ALGORITHM STUDIO'):
        s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(NAVY)
        text(s,.6,.3,12,.35,kicker,11,'#35d9c0',True)
        text(s,.6,.85,12,1,title,32,bold=True)
        text(s,.6,7.1,11,.25,'01AI0506 · Fictional campus · Distances in metres',10,'#91a4bc')
        text(s,12,7.1,.7,.25,f'{len(prs.slides):02d}',10,'#91a4bc')
        return s
    def card(s,x,y,w,h,title,body):
        box=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        box.fill.solid();box.fill.fore_color.rgb=rgb('#142238');box.line.color.rgb=rgb('#2b405c')
        text(s,x+.22,y+.2,w-.44,.55,title,21,'#35d9c0',True)
        text(s,x+.22,y+.9,w-.44,h-.95,body,19)
    def map_graph(s,x,y,w,h,edges=()):
        selected={frozenset(e[:2]) for e in edges}
        def point(n):
            px,py=coordinates[n];return x+.5+px/720*(w-1),y+h-.4-py/570*(h-.8)
        for u,v,weight in graph.edges():
            ux,uy=point(u);vx,vy=point(v)
            line=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(ux),Inches(uy),Inches(vx),Inches(vy))
            line.line.color.rgb=rgb('#35d9c0' if frozenset((u,v)) in selected else '#435775');line.line.width=Pt(3 if frozenset((u,v)) in selected else 1.3)
            text(s,(ux+vx)/2-.2,(uy+vy)/2-.12,.6,.3,str(weight),10,'#bac9dc')
        for n in graph.nodes:
            px,py=point(n)
            node=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(px-.055),Inches(py-.055),Inches(.11),Inches(.11));node.fill.solid();node.fill.fore_color.rgb=rgb('#6b9fff');node.line.fill.background()
            text(s,px-.73,py-.38,1.65,.5,n,11,bold=True)
    s=slide('Smart Campus Navigator','B.TECH AI & ML / DESIGN AND ANALYSIS OF ALGORITHMS')
    text(s,.65,2,6,1.4,'Interactive graph algorithms.\nRoutes students can explain.',29,'#b8cbe4')
    text(s,.65,4.15,5.5,1.9,'[Student Name]  ·  [Enrollment Number]\n[Guide Name]\n[University / Department]',18,'#91a4bc')
    map_graph(s,6.5,1.7,6,4.7,list(zip(route.path,route.path[1:])))
    s.notes_slide.notes_text_frame.text='Introduce yourself. All campus data is fictional. No API key or external maps are required.'
    s=slide('Two campus problems, two objectives')
    card(s,.65,2,5.9,3.8,'A student’s journey','Find the least walking distance between two selected locations.\nExample: Main Gate → Hostel = 830 m.')
    card(s,6.8,2,5.9,3.8,'An administrator’s network','Connect every location with minimum total pathway length.\nExample: 9-edge MST = 1,690 m.')
    text(s,.8,6.1,11.8,.5,'A minimum spanning tree does not guarantee the shortest route for each journey.',20,'#f5c769')
    s=slide('Scope and learning objectives')
    card(s,.65,1.95,3.9,4.6,'Model','10 locations\n17 undirected paths\nNonnegative metres\nFictional coordinates')
    card(s,4.75,1.95,3.9,4.6,'Implement','Six handwritten algorithms\nStep histories\nPath reconstruction\nValidated graph edits')
    card(s,8.85,1.95,3.8,4.6,'Explain','Unit 4 graph / greedy work\nUnit 1 efficiency\nCorrectness and tests\nPractical 8–11 mapping*')
    text(s,.8,6.65,12,.35,'*Mapping supplied in the brief; no separate official syllabus was attached.',12,'#91a4bc')
    s=slide('One graph, six useful perspectives')
    for x,y,title,body in [(0.65,1.95,'BFS / DFS','Explore reachability\nQueue versus stack'),(4.75,1.95,'Dijkstra','One source\nShortest weighted routes'),(8.85,1.95,'Floyd-Warshall','Every pair\nIntermediate-node DP'),(0.65,4.35,'Prim','Grow a minimum tree\nCheapest crossing edge'),(4.75,4.35,'Kruskal','Sort and join\nDSU rejects cycle edges'),(8.85,4.35,'Learning Lab','Fragment-based replay\nLive comparison matrix')]:card(s,x,y,3.85,2.1,title,body)
    s=slide('Readable architecture, local execution')
    for x,title,body in [(0.65,'Dataset','campus.json\nGraph validation'),(3.8,'Algorithms','Standard Python\nResult + events'),(6.95,'Interface','Streamlit controls\nPlotly highlights'),(10.1,'Evidence','unittest checks\nPDF + PPTX')]:card(s,x,2.4,2.65,2.8,title,body)
    text(s,.8,5.75,11.8,.8,'Session edits → recompute → inspect → restore\nNo database, account, API key or cloud subscription.',21,'#b8cbe4')
    s=slide('Dijkstra in action: 830 metres')
    map_graph(s,.4,1.65,8.1,4.9,list(zip(route.path,route.path[1:])))
    text(s,8.8,2.1,3.9,4.3,'Main Gate → Academic Block: 300 m\nAcademic Block → Library: 150 m\nLibrary → Canteen: 180 m\nCanteen → Hostel: 200 m\nTotal: 830 m',18,'#35d9c0')
    text(s,.8,6.5,11.5,.4,'Nearest unsettled first. Relax neighbors. Record parents. Nonnegative weights make settling safe.',16,'#b8cbe4')
    s.notes_slide.notes_text_frame.text='Correct edge sequence: Main Gate–Academic Block 300; Academic Block–Library 150; Library–Canteen 180; Canteen–Hostel 200. Equal 300-m alternatives exist to Academic Block; strict improvements keep the direct parent.'
    s=slide('Explore layers; compare every pair')
    card(s,.65,2,5.9,3.9,'BFS and DFS','BFS: FIFO queue; minimum unweighted hops.\nDFS: stack frames; finish a branch, then backtrack.\nNeither optimizes walking metres.')
    card(s,6.8,2,5.9,3.9,'Floyd-Warshall','Allow one intermediate k at a time.\nd[i,j] = min(d[i,j], d[i,k] + d[k,j])\nNext-hop table reconstructs the route.')
    text(s,.8,6.15,11.5,.6,'Triangle: A–B=2, B–C=1, A–C=5. One hop costs 5; the best weighted route costs 3.',19,'#f5c769')
    s=slide('Minimum network: 1,690 metres')
    map_graph(s,.4,1.6,8.1,5,prim_result.edges)
    card(s,8.8,2.05,3.9,3.8,'9 edges / 10 vertices','Prim grows a connected patch.\nKruskal sorts and joins components.\nSame optimal total; ties can allow different trees.')
    text(s,.8,6.6,11.5,.35,'Main Gate → Auditorium: 380 m shortest route; 510 m along this MST.',17,'#f5c769')
    s=slide('Efficiency: explain the growth')
    rows=[['Algorithm','Time','Space*']]+[[a['name'],a['time'],a['space']] for a in ALGORITHMS]
    table=s.shapes.add_table(len(rows),3,Inches(.75),Inches(1.85),Inches(11.8),Inches(3.8)).table
    for i,row in enumerate(rows):
        for j,value in enumerate(row):
            cell=table.cell(i,j);cell.text=value;cell.fill.solid();cell.fill.fore_color.rgb=rgb('#19364a' if i==0 else '#142238')
            for p in cell.text_frame.paragraphs:p.font.size=Pt(18);p.font.color.rgb=rgb('#e9f0fb')
    text(s,.8,5.95,11.9,.9,'*Without history; excludes graph storage. Lazy heaps retain O(E) entries.\nHistory: O(V+E), except Floyd up to O(V³). Kruskal is usually O(E log E) on connected graphs.',15,'#91a4bc')
    s=slide('Tested answers and measured performance')
    card(s,.65,1.95,4.3,4.6,'Correctness evidence','830 m shortest route\n1,690 m MST\nDisconnected forests\nZero weights and reconstruction\nTiny exhaustive MST oracle')
    # Native editable chart based only on the saved CSV.
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    data=CategoryChartData();data.categories=['12 vertices','24 vertices','48 vertices']
    for name in ['Dijkstra','Floyd-Warshall']:
        data.add_series(name,[float(next(x['Median ms'] for x in bench if x['Algorithm']==name and int(x['V'])==n)) for n in (12,24,48)])
    chart=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(5.25),Inches(2),Inches(7.25),Inches(3.8),data).chart
    chart.has_legend=True;chart.legend.position=XL_LEGEND_POSITION.BOTTOM
    chart.chart_style=10
    text(s,5.45,5.95,6.8,.85,'Recorded median milliseconds; 7 repeats, trace off.\nDifferent workloads: one source versus all pairs. No overall winner implied.',14,'#91a4bc')
    s=slide('Live demonstration: predict, then inspect')
    card(s,.65,2,3.9,4.2,'1 / Route','Dashboard: 10 vertices, 17 edges.\nMain Gate → Hostel: 830 m.\nOpen Dijkstra steps; click Next.')
    card(s,4.75,2,3.9,4.2,'2 / Connect','Compare Prim and Kruskal.\nBoth: 1,690 m, 9 edges.\nExplain route versus network.')
    card(s,8.85,2,3.8,4.2,'3 / Experiment','Add Gate–Hostel at 100 m.\nRecomputed route: 100 m.\nRestore: back to 830 m.')
    text(s,.8,6.4,12,.4,'Then show Floyd’s matrix and BFS/DFS. All controls work locally after installation.',17,'#b8cbe4')
    s=slide('What this project demonstrates')
    text(s,.8,2,7.5,2.5,'A route and a network solve different problems.\nGreedy choices need a correctness argument.\nVisible steps make the code explainable.\nTests and measurements answer different questions.',25,'#e9f0fb')
    card(s,9,2,3.6,3.7,'Next steps','Validated real data\nAccessibility constraints\nReviewed graph versions\nKeep assumptions explicit')
    text(s,.8,5.4,11.7,1.3,'Scope: fictional static distances; symmetric paths; length-based construction cost.\nReferences: Cormen et al., Introduction to Algorithms, 4e (2022); Python and library official docs. Full references in the report.',15,'#91a4bc')
    for s in prs.slides:
        if not s.notes_slide.notes_text_frame.text.strip():s.notes_slide.notes_text_frame.text='Use docs/DEMO_SCRIPT.md for spoken guidance. Source: original campus JSON, implemented algorithms and recorded benchmark CSV.'
    prs.save(str(DOCS/'Project_Presentation.pptx'))
    assert len(prs.slides)==12


if __name__=='__main__':
    generate_study_guides()
    generate_report()
    generate_slides()
    summary={'vertices':len(graph.nodes),'edges':len(graph.edges()),'route':route.path,'distance_m':route.distance,
        'prim_total_m':prim_result.total,'kruskal_total_m':kruskal_result.total,
        'bfs_order':bfs(graph,'Main Gate').order,'dfs_order':dfs(graph,'Main Gate').order,
        'prim_edges':prim_result.edges,'kruskal_edges':kruskal_result.edges}
    (DOCS/'computed_results.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
    print('Generated 12-page PDF, 12-slide PPTX, 42-question viva guide, learning guide and computed results.')
